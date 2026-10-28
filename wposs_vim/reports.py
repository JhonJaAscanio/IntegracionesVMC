"""Exportación a Excel, informe completo de visita y descarga de la carpeta de la visita.

Ninguna de estas operaciones modifica el libro principal.
"""
import io
import posixpath
import tempfile
import zipfile
from datetime import datetime
from pathlib import Path

from flask import render_template

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import schema as S

HEADER_FILL = PatternFill(fill_type="solid", fgColor="FF09919D")
HEADER_FONT = Font(bold=True, color="FFFFFFFF")


def _sheet(wb, title, columns, rows):
    ws = wb.create_sheet(title)
    for i, (_, header) in enumerate(columns, start=1):
        cell = ws.cell(1, i, header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = 18
    for r, rec in enumerate(rows, start=2):
        for i, (key, _) in enumerate(columns, start=1):
            value = rec.get(key, "")
            cell = ws.cell(r, i)
            cell.value = value if value != "" else None
            if isinstance(value, str):
                cell.data_type = "s"
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "B2"
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(rows) + 1}"
    return ws


def export_workbook(service, sessions, folder):
    """Genera un .xlsx con las sesiones filtradas y sus registros relacionados."""
    data = service.data()
    sids = {s["sesion_id"] for s in sessions}
    wb = Workbook()
    wb.remove(wb.active)
    session_cols = S.ID_COLUMNS + [("estado_registro", "Estado de registro"),
                                   ("estado_integracion", "Estado de la integración")]
    _sheet(wb, "Sesiones", session_cols, sessions)
    _sheet(wb, "Revisión física", S.REVISION_COLUMNS, [r for r in data.revisiones if r["sesion_id"] in sids])
    _sheet(wb, "Casos de prueba", S.CASOS_COLUMNS, [t for t in data.pruebas if t["sesion_id"] in sids])
    _sheet(wb, "Incidencias", S.INCIDENCIAS_COLUMNS, [i for i in data.incidencias if i["sesion_id"] in sids])
    totals_cols = [("sesion_id", "Sesión ID"), ("planificadas", "Total planificadas"),
                   ("ejecutadas", "Total ejecutadas"), ("aprobadas", "Aprobadas"),
                   ("fallidas", "Fallidas"), ("bloqueadas", "Bloqueadas"),
                   ("pendientes", "Pendientes"), ("incidencias_abiertas", "Incidencias abiertas"),
                   ("concepto", "Concepto final")]
    totals = []
    for s in sessions:
        t = service.totals(s["sesion_id"], data)
        totals.append({**t, "sesion_id": s["sesion_id"], "concepto": s.get("concepto", "")})
    _sheet(wb, "Totales de cierre", totals_cols, totals)
    folder.mkdir(parents=True, exist_ok=True)
    name = f"Exportacion_WPOSS_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
    path = folder / name
    n = 1
    while path.exists():
        path = folder / f"Exportacion_WPOSS_{datetime.now():%Y%m%d_%H%M%S}_{n}.xlsx"
        n += 1
    wb.save(path)
    return path


def report_context(service, sid, file_link):
    """Datos completos de una visita para el expediente y el informe.

    file_link(ruta) devuelve el enlace a usar para cada evidencia (URL de la app o ruta relativa).
    """
    from .evidence import FOLDER_LABELS, FOLDERS, WEB_IMAGE_EXTENSIONS, folder_of, subfolder_of
    data = service.data()
    bundle = service.session_bundle(sid)
    groups = {k: [] for k in FOLDERS}
    for e in bundle["evidencias"]:
        groups[folder_of(e)].append(e)
    log_groups = {
        "Logs del POS NEW5320": [e for e in groups["logs"] if subfolder_of(e) == "logs/new5320"],
        "Logs de Qibixx": [e for e in groups["logs"] if subfolder_of(e) == "logs/qibixx"],
        "Archivos DEX": [e for e in groups["logs"] if subfolder_of(e) == "logs/dex"],
        "Otros logs de la visita": [e for e in groups["logs"] if subfolder_of(e) == "logs"],
    }
    return {
        "b": bundle, "totals": service.totals(sid, data), "ev": service.evaluate_closure(sid, data),
        "summary": service.email_summary(sid), "groups": groups, "log_groups": log_groups,
        "legacy_fields": service.legacy_fields(bundle["sesion"]),
        "folder_labels": FOLDER_LABELS,
        "image_exts": WEB_IMAGE_EXTENSIONS, "file_link": file_link,
        "incident_history": {i["id"]: service.history("Incidencia", i["id"]) for i in bundle["incidencias"]},
        "generado": datetime.now().strftime("%Y-%m-%d %H:%M"),
    }


def render_visit_report(service, evidence, sid, location):
    """HTML autónomo del informe; los enlaces a evidencias son relativos a `location`.

    location es la carpeta (relativa a data/) donde quedará el HTML, p. ej. evidencias/SES-004/informes.
    """
    def link(ruta):
        return posixpath.relpath(ruta, location)
    return render_template("informe_visita.html", **report_context(service, sid, link))


def generate_visit_report(service, evidence, sid):
    """Genera el informe completo y lo guarda en la carpeta informes/ de la visita."""
    location = f"{evidence.base.name}/{sid}/informes"
    html = render_visit_report(service, evidence, sid, location)
    filename = f"Informe_{sid}_{datetime.now():%Y%m%d_%H%M%S}.html"
    return service.save_report(sid, filename, html)


def build_visit_zip(service, evidence, sid):
    """Carpeta completa de la visita en .zip: archivos, informe actualizado y datos en Excel."""
    folder = evidence.session_dir(sid)
    session = service.get_session(sid)
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as zf:
        if folder.exists():
            for path in sorted(folder.rglob("*")):
                if path.is_file() and not path.name.startswith(("~$", ".")):
                    zf.write(path, f"{sid}/{path.relative_to(folder).as_posix()}")
        html = render_visit_report(service, evidence, sid, f"{evidence.base.name}/{sid}")
        zf.writestr(f"{sid}/informe_visita.html", html)
        with tempfile.TemporaryDirectory() as tmp:
            xlsx = export_workbook(service, [session], Path(tmp))
            zf.write(xlsx, f"{sid}/datos_visita.xlsx")
    buffer.seek(0)
    return buffer
