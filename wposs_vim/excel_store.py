"""Capa única de acceso al libro Excel.

Toda lectura y escritura del libro pasa por ExcelStore:

* La copia de trabajo se crea a partir del original (que nunca se modifica).
* Las escrituras se expresan como operaciones serializables (ver
  operations.py). Si el libro no puede escribirse (por ejemplo, está abierto
  en Excel), las operaciones se guardan en un diario de cambios pendientes
  en disco y se reintentan después; nunca se informa que se guardaron.
* Cada escritura carga el libro, aplica las operaciones, guarda en un archivo
  temporal, lo verifica y reemplaza la copia de trabajo (os.replace).
* Un candado de proceso serializa las escrituras concurrentes.
"""
import hashlib
import io
import json
import logging
import os
import shutil
import threading
import uuid
import warnings
from copy import copy
from datetime import date, datetime, time
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from . import schema as S

log = logging.getLogger("wposs.excel")

# openpyxl avisa que descarta las validaciones de extensión; se regeneran.
warnings.filterwarnings("ignore", message="Data Validation extension is not supported")


class StoreError(Exception):
    """Error genérico de lectura/escritura del libro."""


class ExcelLockedError(StoreError):
    """El libro está abierto o bloqueado por otro programa."""


class PendingWriteError(StoreError):
    """Los cambios no se escribieron; quedaron en el diario de pendientes."""

    def __init__(self, pending_count, reason):
        super().__init__(reason)
        self.pending_count = pending_count
        self.reason = reason


class ConflictError(StoreError):
    """La operación contradice el estado actual del libro (ID duplicado, etc.)."""


# --------------------------------------------------------------------------
# Utilidades de celdas y hojas
# --------------------------------------------------------------------------

def open_workbook(path, read_only=False):
    """Carga el libro desde memoria para no retener el archivo abierto.

    Algunas combinaciones de openpyxl y Python mantienen el .xlsx abierto hasta
    la recolección de basura; en Windows eso impide reemplazarlo después.
    """
    with open(path, "rb") as fh:
        content = fh.read()
    return load_workbook(io.BytesIO(content), read_only=read_only)


def norm_value(value):
    if value is None:
        return ""
    if isinstance(value, datetime):
        if value.hour == 0 and value.minute == 0 and value.second == 0:
            return value.strftime("%Y-%m-%d")
        return value.strftime("%Y-%m-%d %H:%M")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.strftime("%H:%M")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def set_cell(cell, value):
    """Escribe un valor de texto sin que se interprete como fórmula."""
    if value is None or value == "":
        cell.value = None
        return
    if isinstance(value, str):
        cell.value = value
        # Un texto que empieza con "=" se guardaría como fórmula: se fuerza texto.
        cell.data_type = "s"
        return
    cell.value = value


def _norm_header(text):
    return " ".join(str(text or "").split()).casefold()


_THIN = Side(style="thin", color="FFBFBFBF")
DATA_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
DATA_ALIGN = Alignment(wrap_text=True, vertical="top")
NO_FILL = PatternFill(fill_type=None)
CALC_FILL = PatternFill(fill_type="solid", fgColor=S.COLOR_CALC)


class Sheet:
    """Acceso por encabezados a una hoja de datos."""

    def __init__(self, ws, columns):
        self.ws = ws
        self.columns = columns
        headers = {}
        for col in range(1, ws.max_column + 1):
            text = ws.cell(1, col).value
            if text is not None and str(text).strip():
                headers.setdefault(_norm_header(text), col)
        self.col = {}
        for key, header in columns:
            idx = headers.get(_norm_header(header))
            if idx is not None:
                self.col[key] = idx
        self.key = columns[0][0]
        if self.key not in self.col:
            raise StoreError(f"La hoja '{ws.title}' no tiene la columna '{columns[0][1]}'.")

    def missing(self):
        return [header for key, header in self.columns if key not in self.col]

    def rows(self):
        key_col = self.col[self.key]
        for r in range(2, self.ws.max_row + 1):
            if norm_value(self.ws.cell(r, key_col).value) == "":
                continue
            rec = {k: norm_value(self.ws.cell(r, c).value) for k, c in self.col.items()}
            rec["_row"] = r
            yield rec

    def find(self, **criteria):
        for rec in self.rows():
            if all(rec.get(k, "") == v for k, v in criteria.items()):
                return rec
        return None

    def last_used_row(self):
        key_col = self.col[self.key]
        last = 1
        for r in range(2, self.ws.max_row + 1):
            if norm_value(self.ws.cell(r, key_col).value) != "":
                last = r
        return last

    def write(self, row, values, calc_keys=()):
        for key, value in values.items():
            if key not in self.col or key.startswith("_"):
                continue
            set_cell(self.ws.cell(row, self.col[key]), value)
        self.style_row(row, calc_keys)

    def style_row(self, row, calc_keys=()):
        calc_cols = {self.col[k] for k in calc_keys if k in self.col}
        max_col = max(self.col.values())
        for col in range(1, max_col + 1):
            cell = self.ws.cell(row, col)
            cell.border = DATA_BORDER
            cell.alignment = DATA_ALIGN
            cell.fill = copy(CALC_FILL) if col in calc_cols else copy(NO_FILL)

    def append(self, values, calc_keys=()):
        row = self.last_used_row() + 1
        self.write(row, values, calc_keys)
        return row

    def clear_row(self, row):
        for col in self.col.values():
            self.ws.cell(row, col).value = None
        self.style_row(row)


def sheet(wb, name):
    if name not in wb.sheetnames:
        raise StoreError(f"El libro no contiene la hoja '{name}'.")
    return Sheet(wb[name], S.SHEET_COLUMNS[name])


def read_catalogs(wb):
    ws = wb[S.SHEET_LISTAS]
    headers = {}
    for col in range(1, ws.max_column + 1):
        text = ws.cell(1, col).value
        if text:
            headers[_norm_header(text)] = col
    result = {}
    for key, header in S.CATALOGS:
        col = headers.get(_norm_header(header))
        values = []
        if col:
            for r in range(2, ws.max_row + 1):
                v = norm_value(ws.cell(r, col).value)
                if v and v not in values:
                    values.append(v)
        result[key] = values
    return result


def catalog_column(wb, catalog_key):
    ws = wb[S.SHEET_LISTAS]
    header = S.CATALOG_LABELS[catalog_key]
    for col in range(1, ws.max_column + 1):
        if _norm_header(ws.cell(1, col).value) == _norm_header(header):
            return col
    return None


# --------------------------------------------------------------------------
# Fórmulas de la hoja Cierre de Sesión (idénticas a las del libro original)
# --------------------------------------------------------------------------

_ID = "'Identificación y Versiones'"
_LOOKUP_COLS = {
    "fecha_visita": "B", "cliente": "F", "fabricante": "G", "modelo_maquina": "H",
    "serial_maquina": "I", "modelo_pos": "J", "serial_pos": "K", "firmware": "N",
    "mdb_estandar": "O", "mdb_version": "P", "mdb_modulo": "Q",
    "version_financiera": "R", "version_polaris": "S", "participantes": "E",
    "responsable": "D", "objetivo": "L",
}


def cierre_formulas(r):
    a = f"A{r}"
    f = {}
    for key, col in _LOOKUP_COLS.items():
        f[key] = (f'=IF({a}="","",IFERROR(INDEX({_ID}!${col}:${col},'
                  f'MATCH({a},{_ID}!$A:$A,0)),""))')
    cp = "'Casos de Prueba'"
    f["total_planificadas"] = f'=IF({a}="","",COUNTIF({cp}!$B$2:$B$1000,{a}))'
    f["total_ejecutadas"] = (f'=IF({a}="","",COUNTIFS({cp}!$B$2:$B$1000,{a},{cp}!$J$2:$J$1000,'
                             f'"<>Pendiente",{cp}!$J$2:$J$1000,"<>"))')
    for key, estado in (("aprobadas", "Aprobado"), ("fallidas", "Fallido"),
                        ("bloqueadas", "Bloqueado"), ("pendientes", "Pendiente")):
        f[key] = (f'=IF({a}="","",COUNTIFS({cp}!$B$2:$B$1000,{a},'
                  f'{cp}!$J$2:$J$1000,"{estado}"))')
    f["incidencias_abiertas"] = (
        f'=IF({a}="","",COUNTIFS(Incidencias!$C$2:$C$1000,{a},Incidencias!$N$2:$N$1000,"<>Cerrado")'
        f'-COUNTIFS(Incidencias!$C$2:$C$1000,{a},Incidencias!$N$2:$N$1000,""))')
    f["resumen"] = (
        f'=IF(A{r}="","","Resumen de visita — Sesión "&A{r}&CHAR(10)&"Fecha de la visita: "&C{r}'
        f'&CHAR(10)&"Cliente / Comercio: "&D{r}&CHAR(10)&"Máquina: "&E{r}&" "&F{r}&" (Serial: "&G{r}&")"'
        f'&CHAR(10)&"POS: "&H{r}&" (Serial: "&I{r}&") — Firmware: "&J{r}&CHAR(10)&"MDB: "&K{r}&" / "&L{r}'
        f'&" / "&M{r}&CHAR(10)&"Financiera: "&N{r}&"  |  Agente Polaris: "&O{r}&CHAR(10)'
        f'&"Participantes: "&P{r}&" (Responsable: "&Q{r}&")"&CHAR(10)&"Objetivo: "&R{r}&CHAR(10)&CHAR(10)'
        f'&"Resultado de pruebas: "&T{r}&"/"&S{r}&" ejecutadas — Aprobadas: "&U{r}&", Fallidas: "&V{r}'
        f'&", Bloqueadas: "&W{r}&", Pendientes: "&X{r}&CHAR(10)&"Incidencias abiertas: "&Y{r}'
        f'&CHAR(10)&CHAR(10)&"Concepto final: "&AB{r}&CHAR(10)&"Pendientes: "&AA{r})')
    return f


CIERRE_FORMULA_KEYS = list(cierre_formulas(2).keys())


# --------------------------------------------------------------------------
# Migración idempotente de la copia de trabajo
# --------------------------------------------------------------------------

def _copy_header_style(ws, src_col, dst_col):
    src = ws.cell(1, src_col)
    dst = ws.cell(1, dst_col)
    dst.font = copy(src.font)
    dst.fill = copy(src.fill)
    dst.border = copy(src.border)
    dst.alignment = copy(src.alignment)


def _add_columns(ws, columns, added_keys, width=16):
    existing = {_norm_header(ws.cell(1, c).value) for c in range(1, ws.max_column + 1)}
    last = max((c for c in range(1, ws.max_column + 1) if ws.cell(1, c).value), default=1)
    changed = False
    for key, header in columns:
        if key in added_keys and _norm_header(header) not in existing:
            last += 1
            ws.cell(1, last).value = header
            _copy_header_style(ws, 1, last)
            ws.column_dimensions[get_column_letter(last)].width = width
            changed = True
    return changed


def _create_app_sheet(wb, name, header_src):
    ws = wb.create_sheet(name)
    for i, (_, header) in enumerate(S.SHEET_COLUMNS[name], start=1):
        ws.cell(1, i).value = header
        src = header_src.cell(1, 1)
        cell = ws.cell(1, i)
        cell.font = copy(src.font)
        cell.fill = copy(src.fill)
        cell.border = copy(src.border)
        cell.alignment = copy(src.alignment)
        ws.column_dimensions[get_column_letter(i)].width = 18
    ws.freeze_panes = "A2"
    return ws


def _replace_catalog(wb, key, legacy, new):
    """Sustituye un catálogo de Listas solo si conserva exactamente los valores `legacy`.

    Deja constancia en App_Historial. Los registros que ya usan valores anteriores no cambian.
    """
    col = catalog_column(wb, key)
    if col is None:
        return False
    ws = wb[S.SHEET_LISTAS]
    current = [norm_value(ws.cell(r, col).value) for r in range(2, ws.max_row + 1)]
    current = [v for v in current if v]
    if current != list(legacy):
        return False
    for r in range(2, ws.max_row + 1):
        ws.cell(r, col).value = None
    for i, value in enumerate(new, start=2):
        set_cell(ws.cell(i, col), value)
    hist = sheet(wb, S.SHEET_APP_HISTORIAL)
    hist.append({"fecha": datetime.now().strftime("%Y-%m-%d %H:%M:%S"), "tipo_registro": "Catálogo",
                 "registro_id": S.CATALOG_LABELS[key], "campo": "Valores",
                 "anterior": "; ".join(legacy), "nuevo": "; ".join(new),
                 "accion": "Catálogo actualizado por definición de WPOSS"})
    log.info("Catálogo '%s' actualizado: %s", S.CATALOG_LABELS[key], ", ".join(new))
    return True


def read_pairs(wb, key):
    """Pares (valor, país) de un catálogo emparejado en Listas, fila a fila."""
    col = catalog_column(wb, key)
    pair_col = catalog_column(wb, S.CATALOG_PAIRS[key])
    if col is None or pair_col is None:
        return []
    ws = wb[S.SHEET_LISTAS]
    pairs = []
    for r in range(2, ws.max_row + 1):
        value = norm_value(ws.cell(r, col).value)
        if value:
            pairs.append((value, norm_value(ws.cell(r, pair_col).value)))
    return pairs


def ensure_schema(wb):
    """Agrega columnas, catálogos y hojas propias de la aplicación si faltan.

    No modifica celdas de datos existentes. La única sustitución es la de los catálogos
    de países y tipos de prueba definidos por WPOSS (ver `_replace_catalog`), y solo si
    aún conservan los valores anteriores. Devuelve True si hubo cambios.
    """
    for name in S.ORIGINAL_SHEETS:
        if name not in wb.sheetnames:
            raise StoreError(f"El libro no contiene la hoja obligatoria '{name}'.")
    changed = False
    changed |= _add_columns(wb[S.SHEET_ID], S.ID_COLUMNS, S.ID_ADDED)
    changed |= _add_columns(wb[S.SHEET_INCIDENCIAS], S.INCIDENCIAS_COLUMNS, S.INCIDENCIAS_ADDED, 14)

    # Catálogos nuevos en la hoja Listas.
    listas = wb[S.SHEET_LISTAS]
    for key in S.CATALOGS_ADDED:
        if catalog_column(wb, key) is None:
            col = max((c for c in range(1, listas.max_column + 1) if listas.cell(1, c).value), default=0) + 1
            listas.cell(1, col).value = S.CATALOG_LABELS[key]
            if col > 1:
                _copy_header_style(listas, 1, col)
            listas.column_dimensions[get_column_letter(col)].width = 28
            seeds = []
            if key == "pais":
                seeds = list(S.SEED_PAISES)
            elif key == "cliente":
                seeds = [c for c, _ in S.SEED_CLIENTES]
            elif key == "cliente_pais":
                seeds = [p for _, p in S.SEED_CLIENTES]
            elif key == "ciudad":
                seeds = [c for cities in S.SEED_CIUDADES.values() for c in cities]
            elif key == "ciudad_pais":
                seeds = [p for p, cities in S.SEED_CIUDADES.items() for _ in cities]
            elif key == "firmware":
                seeds = list(S.SEED_FIRMWARE)
            elif key == "fabricante":
                ident = sheet(wb, S.SHEET_ID)
                for rec in ident.rows():
                    v = rec.get("fabricante", "")
                    if v and v != S.PENDIENTE_CONFIRMAR and v not in seeds:
                        seeds.append(v)
            elif key == "item_revision":
                rev = sheet(wb, S.SHEET_REVISION)
                for rec in rev.rows():
                    v = rec.get("item", "")
                    if v and v not in seeds:
                        seeds.append(v)
            for i, value in enumerate(seeds, start=2):
                listas.cell(i, col).value = value
            changed = True

    header_src = wb[S.SHEET_ID]
    for name in S.APP_SHEETS:
        if name not in wb.sheetnames:
            _create_app_sheet(wb, name, header_src)
            changed = True
    changed |= _add_columns(wb[S.SHEET_APP_EVIDENCIAS], S.APP_EVIDENCIAS_COLUMNS, S.APP_EVIDENCIAS_ADDED, 18)

    # Listas definidas por WPOSS que sustituyen a las anteriores, solo si siguen sin cambios.
    changed |= _replace_catalog(wb, "pais", S.LEGACY_SEED_PAISES, S.SEED_PAISES)
    changed |= _replace_catalog(wb, "tipo_prueba", S.LEGACY_TIPOS_PRUEBA, S.TIPOS_PRUEBA)

    # Metadatos de sesiones que existían antes de la aplicación.
    meta = sheet(wb, S.SHEET_APP_SESIONES)
    known = {rec["sesion_id"] for rec in meta.rows()}
    cierres = {rec["sesion_id"]: rec for rec in sheet(wb, S.SHEET_CIERRE).rows()}
    now = datetime.now().strftime("%Y-%m-%d %H:%M")
    for rec in sheet(wb, S.SHEET_ID).rows():
        sid = rec["sesion_id"]
        if sid in known:
            continue
        cierre = cierres.get(sid)
        estado = (S.ESTADO_REGISTRO_CERRADA if cierre and cierre.get("concepto")
                  else S.ESTADO_REGISTRO_EN_CURSO)
        meta.append({"sesion_id": sid, "estado_registro": estado, "creado": now,
                     "modificado": now, "origen": "Libro original"})
        known.add(sid)
        changed = True
    return changed


def apply_validations(wb):
    """Regenera las listas desplegables como validaciones estándar de Excel."""
    catalogs = read_catalogs(wb)
    for ws in wb.worksheets:
        ws.data_validations.dataValidation = []
    for sheet_name, field, catalog in S.DATA_VALIDATIONS:
        sh = sheet(wb, sheet_name)
        col_idx = sh.col.get(field)
        cat_col = catalog_column(wb, catalog)
        if not col_idx or not cat_col:
            continue
        count = max(len(catalogs.get(catalog, [])), 1)
        letter = get_column_letter(cat_col)
        dv = DataValidation(type="list", allow_blank=True,
                            formula1=f"{S.SHEET_LISTAS}!${letter}$2:${letter}${count + 1}")
        dv.showErrorMessage = False
        target = get_column_letter(col_idx)
        dv.add(f"{target}2:{target}{S.DV_MAX_ROW}")
        sh.ws.add_data_validation(dv)


def file_sha256(path):
    digest = hashlib.sha256()
    with open(path, "rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 16), b""):
            digest.update(chunk)
    return digest.hexdigest()


def validate_workbook_file(path):
    """Comprueba que un archivo sea un libro WPOSS válido. Devuelve lista de problemas."""
    problems = []
    try:
        wb = open_workbook(path)
    except Exception as exc:  # archivo dañado o no es xlsx
        return [f"No se pudo abrir el archivo como libro Excel: {exc}"]
    try:
        for name in S.ORIGINAL_SHEETS:
            if name not in wb.sheetnames:
                problems.append(f"Falta la hoja '{name}'.")
        if problems:
            return problems
        for name in (S.SHEET_ID, S.SHEET_REVISION, S.SHEET_CASOS, S.SHEET_INCIDENCIAS, S.SHEET_CIERRE):
            sh = Sheet(wb[name], S.SHEET_COLUMNS[name])
            added = set(S.ID_ADDED) | set(S.INCIDENCIAS_ADDED)
            missing = [h for k, h in sh.columns if k not in sh.col and k not in added]
            if missing:
                problems.append(f"La hoja '{name}' no tiene las columnas: {', '.join(missing)}.")
        ids = [rec["sesion_id"] for rec in Sheet(wb[S.SHEET_ID], S.ID_COLUMNS).rows()]
        dups = sorted({i for i in ids if ids.count(i) > 1})
        if dups:
            problems.append(f"Sesión ID duplicado: {', '.join(dups)}.")
    except StoreError as exc:
        problems.append(str(exc))
    finally:
        wb.close()
    return problems


# --------------------------------------------------------------------------
# Almacén
# --------------------------------------------------------------------------

class Data:
    """Instantánea de solo lectura del libro."""

    def __init__(self, wb):
        self.sesiones = list(sheet(wb, S.SHEET_ID).rows())
        self.revisiones = list(sheet(wb, S.SHEET_REVISION).rows())
        self.pruebas = list(sheet(wb, S.SHEET_CASOS).rows())
        self.incidencias = list(sheet(wb, S.SHEET_INCIDENCIAS).rows())
        self.cierres = list(sheet(wb, S.SHEET_CIERRE).rows())
        self.meta = {r["sesion_id"]: r for r in sheet(wb, S.SHEET_APP_SESIONES).rows()}
        self.evidencias = list(sheet(wb, S.SHEET_APP_EVIDENCIAS).rows())
        self.historial = list(sheet(wb, S.SHEET_APP_HISTORIAL).rows())
        self.contadores = {r["prefijo"]: r["ultimo"] for r in sheet(wb, S.SHEET_APP_CONTADORES).rows()}
        self.catalogos = read_catalogs(wb)
        self.clientes_pais = dict(read_pairs(wb, "cliente"))
        self.ciudades = {}
        for ciudad, pais in read_pairs(wb, "ciudad"):
            self.ciudades.setdefault(pais, []).append(ciudad)


def excel_lock_present(path):
    """Excel crea '~$<nombre>' junto al libro mientras lo tiene abierto."""
    path = Path(path)
    suffix = path.name[2:] if len(path.name) > 2 else path.name
    for candidate in path.parent.glob("~$*"):
        if candidate.name.endswith(suffix):
            return True
    return False


class ExcelStore:
    def __init__(self, paths, operations, before_write=None):
        self.paths = paths
        self.operations = operations
        self.before_write = before_write
        self._lock = threading.RLock()
        self._cache = None
        self._cache_key = None

    # -- inicialización ----------------------------------------------------
    def initialize(self):
        """Crea la copia de trabajo a partir del original si no existe."""
        paths = self.paths
        paths.ensure()
        for stale in paths.trabajo.parent.glob(".tmp_*.xlsx"):
            try:
                stale.unlink()
            except OSError:
                pass
        if not paths.trabajo.exists():
            if not paths.original.exists():
                raise StoreError(f"No se encontró el libro original en: {paths.original}")
            problems = validate_workbook_file(paths.original)
            if problems:
                raise StoreError("El libro original no es válido: " + " ".join(problems))
            if not paths.original_copia.exists():
                shutil.copy2(paths.original, paths.original_copia)
                with open(paths.original_copia.parent / "ORIGEN.txt", "w", encoding="utf-8") as fh:
                    fh.write(f"Copia íntegra del libro original, tomada el {datetime.now():%Y-%m-%d %H:%M:%S}\n"
                             f"Origen: {paths.original}\nSHA-256: {file_sha256(paths.original)}\n")
            wb = open_workbook(paths.original)
            ensure_schema(wb)
            self._save_atomic(wb)
            log.info("Copia de trabajo creada desde el original: %s", paths.trabajo)
        else:
            with self._lock:
                wb = open_workbook(paths.trabajo)
                if ensure_schema(wb):
                    if self.before_write:
                        self.before_write()
                    self._save_atomic(wb)
                    log.info("Esquema de la copia de trabajo actualizado.")

    # -- lectura -----------------------------------------------------------
    def read(self):
        path = self.paths.trabajo
        try:
            st = path.stat()
        except FileNotFoundError:
            raise StoreError("No existe la copia de trabajo del libro.")
        key = (st.st_mtime_ns, st.st_size)
        if self._cache is not None and self._cache_key == key:
            return self._cache
        with self._lock:
            try:
                wb = open_workbook(path)
            except PermissionError:
                raise ExcelLockedError("No se pudo leer el libro: está bloqueado por otro programa.")
            except Exception as exc:
                raise StoreError(f"No se pudo leer el libro: {exc}")
            try:
                data = Data(wb)
            finally:
                wb.close()
            self._cache, self._cache_key = data, key
            return data

    def invalidate(self):
        self._cache = None
        self._cache_key = None

    # -- diario de cambios pendientes -------------------------------------
    def pending(self):
        items = []
        for path in sorted(self.paths.pendientes.glob("*.json")):
            try:
                with open(path, "r", encoding="utf-8") as fh:
                    item = json.load(fh)
                item["_file"] = str(path)
                items.append(item)
            except (OSError, ValueError):
                log.error("Archivo de cambios pendientes ilegible: %s", path)
        return items

    def _journal(self, ops, descripcion):
        """Guarda un cambio (una o varias operaciones) como una sola entrada."""
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
        item = {"ops": [[name, payload] for name, payload in ops], "descripcion": descripcion,
                "creado": datetime.now().strftime("%Y-%m-%d %H:%M:%S")}
        target = self.paths.pendientes / f"{stamp}_{uuid.uuid4().hex[:6]}.json"
        tmp = target.with_suffix(".tmp")
        with open(tmp, "w", encoding="utf-8") as fh:
            json.dump(item, fh, ensure_ascii=False, indent=2)
        os.replace(tmp, target)

    # -- escritura ---------------------------------------------------------
    def write(self, ops, descripcion):
        """Aplica operaciones [(nombre, payload)]. Primero reaplica pendientes."""
        with self._lock:
            pending = self.pending()
            try:
                self._commit(pending, ops)
            except ExcelLockedError as exc:
                self._journal(ops, descripcion)
                count = len(pending) + 1
                log.warning("Escritura diferida (%s): %s", descripcion, exc)
                raise PendingWriteError(count, str(exc))
            for item in pending:
                try:
                    os.remove(item["_file"])
                except OSError:
                    pass
            log.info("Guardado: %s", descripcion)

    def retry_pending(self):
        with self._lock:
            pending = self.pending()
            if not pending:
                return 0
            self._commit(pending, [])
            for item in pending:
                try:
                    os.remove(item["_file"])
                except OSError:
                    pass
            log.info("Se aplicaron %d cambios pendientes.", len(pending))
            return len(pending)

    def _check_writable(self):
        path = self.paths.trabajo
        if excel_lock_present(path):
            raise ExcelLockedError("El libro está abierto en Excel. Ciérrelo para guardar los cambios.")
        try:
            with open(path, "r+b"):
                pass
        except PermissionError:
            raise ExcelLockedError("El libro está bloqueado por otro programa (posiblemente Excel).")

    def _commit(self, pending, ops):
        self._check_writable()
        if self.before_write:
            self.before_write()
        wb = open_workbook(self.paths.trabajo)
        ensure_schema(wb)
        failed_dir = self.paths.pendientes / "fallidos"
        for item in pending:
            try:
                for name, payload in item["ops"]:
                    self.operations.apply(wb, name, payload)
            except (StoreError, KeyError, ValueError, TypeError) as exc:
                failed_dir.mkdir(exist_ok=True)
                shutil.move(item["_file"], failed_dir / Path(item["_file"]).name)
                item["_file"] = ""
                log.error("Cambio pendiente descartado (%s): %s", item.get("descripcion"), exc)
        for name, payload in ops:
            self.operations.apply(wb, name, payload)
        self._save_atomic(wb)

    def _save_atomic(self, wb):
        apply_validations(wb)
        wb.calculation.fullCalcOnLoad = True
        target = self.paths.trabajo
        target.parent.mkdir(parents=True, exist_ok=True)
        tmp = target.parent / f".tmp_{uuid.uuid4().hex}.xlsx"
        try:
            wb.save(tmp)
            check = open_workbook(tmp, read_only=True)
            names = check.sheetnames
            check.close()
            missing = [n for n in S.ORIGINAL_SHEETS + S.APP_SHEETS if n not in names]
            if missing:
                raise StoreError(f"Verificación fallida, faltan hojas: {missing}")
            try:
                os.replace(tmp, target)
            except PermissionError:
                raise ExcelLockedError("El libro está bloqueado por otro programa (posiblemente Excel).")
        finally:
            if tmp.exists():
                try:
                    tmp.unlink()
                except OSError:
                    pass
        self.invalidate()
