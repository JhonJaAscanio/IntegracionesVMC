"""Integridad del libro, bloqueo, respaldos, evidencias, exportación y dashboard vacío."""
import io

import pytest
from openpyxl import load_workbook

from conftest import SESSION_FORM, Client, make_app, prueba_data, service
from wposs_vim import schema as S
from wposs_vim.excel_store import PendingWriteError, StoreError, file_sha256


def _snapshot(path, sheets):
    wb = load_workbook(path)
    snap = {}
    for name in sheets:
        ws = wb[name]
        snap[name] = {c.coordinate: c.value for row in ws.iter_rows() for c in row if c.value is not None}
    return snap


def test_11_hojas_y_formulas_originales_conservadas(app):
    svc = service(app)
    original = app.extensions["wposs"]["paths"].original
    original_hash = file_sha256(original)
    work = app.extensions["wposs"]["paths"].trabajo
    before = _snapshot(original, S.ORIGINAL_SHEETS)
    sid = svc.create_session(SESSION_FORM)
    svc.create_test(prueba_data(sid))
    svc.init_revision(sid)
    svc.save_cierre(sid, {"fecha_cierre": "2026-10-28", "concepto": "Pendiente", "responsable_cierre": "QA"})
    after = _snapshot(work, S.ORIGINAL_SHEETS)
    # Instrucciones intacta.
    assert after[S.SHEET_INSTRUCCIONES] == before[S.SHEET_INSTRUCCIONES]
    # Todas las celdas originales de las hojas de datos siguen igual (solo se agregan filas/columnas).
    for name in S.ORIGINAL_SHEETS:
        for coord, value in before[name].items():
            if name == S.SHEET_CIERRE and coord in ("A5", "B5", "Z5", "AA5", "AB5", "AC5"):
                continue  # fila 5 preparada que recibió el nuevo cierre
            if name == S.SHEET_LISTAS and coord[0] == "E" and coord != "E1":
                continue  # tipos de prueba reemplazados por definición de WPOSS (se verifica abajo)
            assert after[name].get(coord) == value, (name, coord)
    tipos = [after[S.SHEET_LISTAS].get(f"E{r}") for r in range(2, 2 + len(S.TIPOS_PRUEBA))]
    assert tipos == S.TIPOS_PRUEBA and "E6" not in after[S.SHEET_LISTAS]
    # Fórmulas del cierre presentes en todas las filas preparadas.
    wb = load_workbook(work)
    ws = wb[S.SHEET_CIERRE]
    for r in range(2, 20):
        assert str(ws.cell(r, 3).value).startswith("=IF(")
        assert str(ws.cell(r, 30).value).startswith("=IF(")
    # Listas desplegables regeneradas (openpyxl elimina las originales x14).
    dvs = {str(dv.sqref): dv.formula1 for dv in wb[S.SHEET_CASOS].data_validations.dataValidation}
    assert any("Listas!$A$2" in f for f in dvs.values())
    assert wb[S.SHEET_LISTAS].sheet_state == "hidden"
    # El archivo original nunca se modifica.
    assert file_sha256(original) == original_hash
    assert file_sha256(app.extensions["wposs"]["paths"].original_copia) == original_hash


def test_12_excel_bloqueado_conserva_cambios_pendientes(app):
    paths = app.extensions["wposs"]["paths"]
    store = app.extensions["wposs"]["store"]
    lock = paths.trabajo.parent / ("~$" + paths.trabajo.name[2:])
    lock.write_text("lock")
    c = Client(app)
    r = c.post("/visitas/nueva", {**SESSION_FORM, "accion": "guardar"})
    assert r.status_code == 302
    page = c.get(r.headers["Location"]).get_data(as_text=True)
    assert "NO se guardó" in page
    assert len(store.pending()) == 1
    assert service(app).get_session("SES-004") is None  # no se afirma un guardado inexistente
    # El ID reservado no se reutiliza mientras está pendiente.
    assert service(app).next_id("SES") == "SES-005"
    lock.unlink()
    assert store.retry_pending() == 1
    assert service(app).get_session("SES-004")["comercio"] == "Comercio Piloto QA"
    assert store.pending() == []


def test_12b_archivo_abierto_por_otro_proceso(app):
    """Un manejador abierto sin permiso de borrado (como Excel) impide reemplazar el archivo."""
    import os
    paths = app.extensions["wposs"]["paths"]
    store = app.extensions["wposs"]["store"]
    if os.name != "nt":
        pytest.skip("Bloqueo de archivos propio de Windows")
    with open(paths.trabajo, "rb"):
        with pytest.raises(PendingWriteError):
            service(app).create_session(SESSION_FORM)
    assert len(store.pending()) == 1
    # La siguiente escritura aplica primero lo pendiente, en orden.
    sid = service(app).create_session(SESSION_FORM)
    assert sid == "SES-005"
    ids = [s["sesion_id"] for s in service(app).data().sesiones]
    assert "SES-004" in ids and "SES-005" in ids and store.pending() == []


def test_13_respaldo_y_restauracion(client, app):
    backups = app.extensions["wposs"]["backups"]
    store = app.extensions["wposs"]["store"]
    name = backups.create("manual")
    sid = service(app).create_session(SESSION_FORM)
    assert service(app).get_session(sid) is not None
    r = client.post("/respaldos/restaurar", {"nombre": name, "confirmar": "si"})
    assert r.status_code == 302
    assert service(app).get_session(sid) is None
    names = [b["nombre"] for b in backups.list()]
    assert any(n.startswith("previo_restauracion_") for n in names)
    # El estado previo a la restauración se puede recuperar.
    prev = next(n for n in names if n.startswith("previo_restauracion_"))
    backups.restore(prev, store)
    assert service(app).get_session(sid) is not None


def test_13b_restaurar_valida_archivo(app):
    backups = app.extensions["wposs"]["backups"]
    store = app.extensions["wposs"]["store"]
    bad = backups.paths.respaldos / "manual_20260101_000000.xlsx"
    bad.write_bytes(b"no es un excel")
    with pytest.raises(StoreError, match="no es válido"):
        backups.restore(bad.name, store)
    with pytest.raises(StoreError):
        backups.restore("../trabajo/WPOSS_Vending_Datos.xlsx", store)
    assert service(app).get_session("SES-002") is not None


def test_13c_retencion_de_respaldos(workdir):
    app = make_app(workdir, respaldos={"retener": 2})
    backups = app.extensions["wposs"]["backups"]
    for _ in range(4):
        backups.create("auto")
    backups.create("manual")
    kinds = [b["tipo"] for b in backups.list()]
    assert kinds.count("auto") == 2 and kinds.count("manual") == 1


def _upload(c, sid, filename, content, tipo="Visita", registro=""):
    return c.post(f"/visitas/{sid}/evidencias", {"tipo_registro": tipo, "registro_id": registro,
                                                 "archivo": (io.BytesIO(content), filename)},
                  content_type="multipart/form-data")


def test_14_validacion_de_adjuntos(app):
    c = Client(app)
    svc = service(app)
    sid = svc.create_session(SESSION_FORM)
    tid = svc.create_test(prueba_data(sid))
    folder = app.extensions["wposs"]["paths"].evidencias / sid
    # Extensión no permitida.
    _upload(c, sid, "malware.exe", b"MZ...")
    # Vacío.
    _upload(c, sid, "vacio.txt", b"")
    assert not [p for p in folder.rglob("*") if p.is_file()]
    # Nombre manipulado: se normaliza dentro de la carpeta de la sesión.
    _upload(c, sid, "../../config.json", b'{"a": 1}', "Prueba", tid)
    _upload(c, sid, "foto conexión.jpg", b"\xff\xd8\xff data")
    _upload(c, sid, "foto conexión.jpg", b"\xff\xd8\xff other")
    files = sorted(p.relative_to(folder).as_posix() for p in folder.rglob("*") if p.is_file())
    assert f"logs/{tid}_config.json" in files
    # Sin sobrescritura: el segundo archivo con el mismo nombre recibe un sufijo.
    assert "fotos/maquina_foto_conexion.jpg" in files and "fotos/maquina_foto_conexion-1.jpg" in files
    data = svc.data()
    evs = [e for e in data.evidencias if e["sesion_id"] == sid]
    assert len(evs) == 3 and all(e["ruta"].startswith(f"evidencias/{sid}/") for e in evs)
    test = next(t for t in data.pruebas if t["id"] == tid)
    assert test["evidencia"] == f"evidencias/{sid}/logs/{tid}_config.json"
    # Recuperable por la ruta registrada.
    assert c.get(f"/archivo/evidencias/{sid}/no_registrado.jpg").status_code == 404
    r = c.get(f"/archivo/{evs[1]['ruta']}")
    assert r.status_code == 200 and r.data == b"\xff\xd8\xff data"


def test_14b_tamano_maximo_y_aviso_tarjeta(workdir):
    app = make_app(workdir, evidencias={"tamano_maximo_mb": 1})
    c = Client(app)
    sid = service(app).create_session(SESSION_FORM)
    r = _upload(c, sid, "grande.log", b"0" * (1024 * 1024 + 10))
    assert service(app).data().evidencias == []
    r = _upload(c, sid, "trx.log", b"PAN=4111 1111 1111 1111 OK\n")
    page = c.get(r.headers["Location"]).get_data(as_text=True)
    assert "formato de número de tarjeta" in page
    ev = service(app).data().evidencias[0]
    assert "tarjeta" in ev["advertencia"]


def test_exportacion_no_modifica_libro(client, app):
    work = app.extensions["wposs"]["paths"].trabajo
    h = file_sha256(work)
    r = client.post("/reportes/exportar", {"fabricante": "Fabricante B"})
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.data))
    assert wb.sheetnames == ["Sesiones", "Revisión física", "Casos de prueba", "Incidencias", "Totales de cierre"]
    assert wb["Sesiones"].max_row == 2 and wb["Sesiones"]["A2"].value == "SES-003"
    assert file_sha256(work) == h
    r2 = client.post("/reportes/exportar", {})
    assert r2.status_code == 200
    assert len(list(app.extensions["wposs"]["paths"].exportaciones.glob("*.xlsx"))) == 2


def test_16_dashboard_sin_datos_ficticios(empty_workdir):
    app = make_app(empty_workdir)
    c = Client(app)
    html = c.get("/").get_data(as_text=True)
    assert "Aún no hay sesiones registradas" in html
    assert 'class="kpi-value' not in html
    stats = service(app).dashboard({})
    assert stats["total_sesiones"] == 0 and stats["pruebas_total"] == 0
    assert stats["inc_abiertas"] == 0 and stats["recientes"] == []
    assert service(app).next_id("SES") == "SES-001"


def test_16b_dashboard_refleja_datos_reales(app):
    stats = service(app).dashboard({})
    # SES-001 es fila de ejemplo y se excluye; SES-002 y SES-003 son reales.
    assert stats["total_sesiones"] == 2
    data = service(app).data()
    real = [t for t in data.pruebas if t["sesion_id"] in ("SES-002", "SES-003")]
    assert stats["pruebas_total"] == len(real)
    assert stats["inc_abiertas"] == 0  # INC-002/INC-003 son filas de ejemplo
