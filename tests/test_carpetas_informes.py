"""Evidencias al crear la visita, carpetas por visita, informes y descarga de la carpeta."""
import io
import zipfile

from openpyxl import load_workbook

from conftest import SESSION_FORM, Client, make_app, prueba_data, service
from wposs_vim import schema as S


def _new_visit(c, files, **extra):
    data = {**SESSION_FORM, "accion": "guardar", **extra, **files}
    return c.post("/visitas/nueva", data, content_type="multipart/form-data")


def _files(app, sid):
    folder = app.extensions["wposs"]["paths"].evidencias / sid
    return sorted(p.relative_to(folder).as_posix() for p in folder.rglob("*") if p.is_file())


def test_nueva_visita_con_fotos_y_logs(app):
    c = Client(app)
    r = _new_visit(c, {
        "fotos_maquina": [(io.BytesIO(b"\xff\xd8\xffA"), "frente.jpg"), (io.BytesIO(b"\x89PNG"), "lado.png")],
        "fotos_etiqueta": (io.BytesIO(b"\xff\xd8\xffB"), "placa serial.jpg"),
        "logs": (io.BytesIO(b"POLL\nACK\n"), "sniffer.log"),
    }, descripcion_evidencias="Visita inicial")
    assert r.status_code == 302
    sid = r.headers["Location"].rsplit("/", 1)[-1]
    assert _files(app, sid) == ["fotos/etiqueta_placa_serial.jpg", "fotos/maquina_frente.jpg",
                                "fotos/maquina_lado.png", "logs/sniffer.log"]
    folder = app.extensions["wposs"]["paths"].evidencias / sid
    assert all((folder / sub).is_dir() for sub in ("fotos", "logs", "videos", "documentos", "informes", "otros"))
    evs = [e for e in service(app).data().evidencias if e["sesion_id"] == sid]
    assert len(evs) == 4
    assert len({e["id"] for e in evs}) == 4  # IDs consecutivos y únicos en una sola escritura
    assert {e["categoria"] for e in evs} == {"Foto de la máquina", "Foto de etiqueta", "Log de la visita"}
    assert all(e["descripcion"] == "Visita inicial" and e["tipo_registro"] == "Visita" for e in evs)


def test_nueva_visita_con_archivo_invalido_no_crea_nada(app):
    c = Client(app)
    before = len(service(app).data().sesiones)
    r = _new_visit(c, {"fotos_maquina": (io.BytesIO(b"texto"), "no_es_foto.txt")})
    assert r.status_code == 422
    html = r.get_data(as_text=True)
    assert "no corresponde a la categoría" in html and "selecciónelos de nuevo" in html
    assert 'value="Comercio Piloto QA"' in html  # datos del formulario conservados
    assert len(service(app).data().sesiones) == before
    assert not (app.extensions["wposs"]["paths"].evidencias / "SES-004").exists()


def test_columna_categoria_agregada_al_libro(app):
    wb = load_workbook(app.extensions["wposs"]["paths"].trabajo)
    headers = [c.value for c in wb[S.SHEET_APP_EVIDENCIAS][1]]
    assert headers[-1] == "Categoría"


def test_carpetas_y_vista_de_carpeta(app):
    c = Client(app)
    sid = service(app).create_session(SESSION_FORM)
    c.post(f"/visitas/{sid}/evidencias", {"tipo_registro": "Visita", "categoria": "Foto de etiqueta",
                                         "archivo": [(io.BytesIO(b"\xff\xd8\xffX"), "etq1.jpg"),
                                                     (io.BytesIO(b"\xff\xd8\xffY"), "etq2.jpg")]},
           content_type="multipart/form-data")
    lista = c.get("/carpetas").get_data(as_text=True)
    assert sid in lista
    page = c.get(f"/carpetas/{sid}?sub=fotos").get_data(as_text=True)
    assert "etiqueta_etq1.jpg" in page and "etiqueta_etq2.jpg" in page
    datos = c.get(f"/carpetas/{sid}?vista=expediente").get_data(as_text=True)
    for texto in ("Comercio Piloto QA", "DATAFAST", "Fabricante QA", "POS-SERIAL-1", "13.0.22", "Resumen de visita"):
        assert texto in datos
    assert c.get("/carpetas/SES-999").status_code == 404
    assert c.get(f"/carpetas/{sid}?sub=../../").status_code == 404


def test_ver_log_en_pagina(app):
    c = Client(app)
    sid = service(app).create_session(SESSION_FORM)
    c.post(f"/visitas/{sid}/evidencias", {"tipo_registro": "Visita",
                                         "archivo": (io.BytesIO(b"<b>POLL</b> 0x10\n"), "bus.log")},
           content_type="multipart/form-data")
    ev = next(e for e in service(app).data().evidencias if e["sesion_id"] == sid)
    r = c.get(f"/carpetas/{sid}/ver?ruta={ev['ruta']}")
    assert r.status_code == 200
    assert "&lt;b&gt;POLL&lt;/b&gt;" in r.get_data(as_text=True)  # escapado, no interpretado
    assert c.get(f"/carpetas/{sid}/ver?ruta=evidencias/{sid}/logs/otro.log").status_code == 404


def test_generar_informe_completo(app):
    c = Client(app)
    svc = service(app)
    sid = svc.create_session(SESSION_FORM)
    svc.create_test(prueba_data(sid, estado="Fallido", resultado_obtenido="Sin respuesta al POLL"))
    c.post(f"/visitas/{sid}/evidencias", {"tipo_registro": "Visita", "categoria": "Foto de la máquina",
                                         "archivo": (io.BytesIO(b"\xff\xd8\xffZ"), "frente.jpg")},
           content_type="multipart/form-data")
    r = c.post(f"/carpetas/{sid}/informe")
    assert r.status_code == 302
    informe = next(e for e in svc.data().evidencias if e.get("categoria") == "Informe generado")
    assert informe["ruta"].startswith(f"evidencias/{sid}/informes/Informe_{sid}_")
    html = c.get(f"/archivo/{informe['ruta']}")
    assert html.status_code == 200 and html.mimetype == "text/html"
    text = html.get_data(as_text=True)
    for texto in ("Informe de visita " + sid, "DATAFAST - Comercio Piloto QA", "Sin respuesta al POLL",
                  'src="../fotos/maquina_frente.jpg"', "Resumen de visita"):
        assert texto in text
    assert "<script" not in text
    # Desde la sección de informes también se genera.
    assert c.post("/reportes/informe", {"sid": sid}).status_code == 302
    assert len([e for e in svc.data().evidencias if e.get("categoria") == "Informe generado"]) == 2
    assert "Informes de visita" in c.get("/reportes").get_data(as_text=True)


def test_descargar_carpeta_zip(app):
    c = Client(app)
    svc = service(app)
    sid = svc.create_session(SESSION_FORM)
    c.post(f"/visitas/{sid}/evidencias", {"tipo_registro": "Visita",
                                         "archivo": (io.BytesIO(b"LOG"), "diag.txt")},
           content_type="multipart/form-data")
    work = app.extensions["wposs"]["paths"].trabajo
    before = work.read_bytes()
    r = c.get(f"/carpetas/{sid}/descargar")
    assert r.status_code == 200 and r.mimetype == "application/zip"
    zf = zipfile.ZipFile(io.BytesIO(r.data))
    names = set(zf.namelist())
    assert {f"{sid}/logs/diag.txt", f"{sid}/informe_visita.html", f"{sid}/datos_visita.xlsx"} <= names
    datos = load_workbook(io.BytesIO(zf.read(f"{sid}/datos_visita.xlsx")))
    assert datos["Sesiones"]["A2"].value == sid
    assert work.read_bytes() == before  # el libro principal no cambia


def test_registrar_archivo_copiado_manualmente(app):
    c = Client(app)
    sid = service(app).create_session(SESSION_FORM)
    folder = app.extensions["wposs"]["paths"].evidencias / sid
    (folder / "fotos" / "copiada.jpg").write_bytes(b"\xff\xd8\xffM")
    page = c.get(f"/carpetas/{sid}").get_data(as_text=True)
    assert "Archivos copiados manualmente" in page and "copiada.jpg" in page
    r = c.post(f"/carpetas/{sid}/registrar", {"ruta": f"evidencias/{sid}/fotos/copiada.jpg",
                                               "categoria": "Foto de etiqueta"})
    assert r.status_code == 302
    ev = next(e for e in service(app).data().evidencias if e["archivo"] == "copiada.jpg")
    assert ev["categoria"] == "Foto de etiqueta" and ev["sesion_id"] == sid
    # Fuera de la carpeta de la visita no se permite registrar.
    c.post(f"/carpetas/{sid}/registrar", {"ruta": "trabajo/WPOSS_Vending_Datos.xlsx"})
    assert not any(e["archivo"] == "WPOSS_Vending_Datos.xlsx" for e in service(app).data().evidencias)


def test_carpeta_vacia_y_libro_bloqueado_conserva_archivos(workdir):
    app = make_app(workdir)
    c = Client(app)
    paths = app.extensions["wposs"]["paths"]
    lock = paths.trabajo.parent / ("~$" + paths.trabajo.name[2:])
    lock.write_text("lock")
    r = _new_visit(c, {"fotos_maquina": (io.BytesIO(b"\xff\xd8\xffL"), "f.jpg")})
    assert r.status_code == 302
    assert service(app).get_session("SES-004") is None
    assert (paths.evidencias / "SES-004" / "fotos" / "maquina_f.jpg").exists()  # el archivo se conserva
    lock.unlink()
    app.extensions["wposs"]["store"].retry_pending()
    evs = [e for e in service(app).data().evidencias if e["sesion_id"] == "SES-004"]
    assert service(app).get_session("SES-004") is not None and len(evs) == 1


def test_logs_new5320_y_qibixx_en_subcarpetas(app):
    c = Client(app)
    r = _new_visit(c, {
        "logs_new5320": [(io.BytesIO(b"MDB RX 0x10"), "pos.log"), (io.BytesIO(b"APP"), "app.txt")],
        "logs_qibixx": (io.BytesIO(b"QIBIXX VEND"), "qibixx.log"),
        "logs": (io.BytesIO(b"SNIFFER"), "sniffer.log"),
    })
    sid = r.headers["Location"].rsplit("/", 1)[-1]
    assert _files(app, sid) == ["logs/new5320/app.txt", "logs/new5320/pos.log",
                                "logs/qibixx/qibixx.log", "logs/sniffer.log"]
    folder = app.extensions["wposs"]["paths"].evidencias / sid
    assert (folder / "logs" / "new5320").is_dir() and (folder / "logs" / "qibixx").is_dir()
    cats = {e["archivo"]: e["categoria"] for e in service(app).data().evidencias if e["sesion_id"] == sid}
    assert cats == {"pos.log": "Log del POS NEW5320", "app.txt": "Log del POS NEW5320",
                    "qibixx.log": "Log de Qibixx", "sniffer.log": "Log de la visita"}
    # Vista de la subcarpeta: solo sus archivos.
    page = c.get(f"/carpetas/{sid}?sub=logs/qibixx").get_data(as_text=True)
    assert "qibixx.log" in page and "pos.log" not in page and "sniffer.log" not in page
    todos = c.get(f"/carpetas/{sid}?sub=logs").get_data(as_text=True)
    assert all(n in todos for n in ("qibixx.log", "pos.log", "sniffer.log"))
    assert c.get(f"/carpetas/{sid}?sub=logs/otro").status_code == 404
    # El informe separa los logs por origen.
    c.post(f"/carpetas/{sid}/informe")
    informe = next(e for e in service(app).data().evidencias if e.get("categoria") == "Informe generado")
    html = c.get(f"/archivo/{informe['ruta']}").get_data(as_text=True)
    assert "Logs del POS NEW5320 (2)" in html and "Logs de Qibixx (1)" in html
    assert "Otros logs de la visita (1)" in html and 'href="../logs/qibixx/qibixx.log"' in html


def test_log_copiado_en_subcarpeta_sugiere_categoria(app):
    c = Client(app)
    sid = service(app).create_session(SESSION_FORM)
    folder = app.extensions["wposs"]["paths"].evidencias / sid
    (folder / "logs" / "qibixx" / "manual.log").write_bytes(b"QIBIXX")
    page = c.get(f"/carpetas/{sid}").get_data(as_text=True)
    assert "logs/qibixx/manual.log" in page
    assert "<option selected>Log de Qibixx</option>" in page
    c.post(f"/carpetas/{sid}/registrar", {"ruta": f"evidencias/{sid}/logs/qibixx/manual.log"})
    ev = next(e for e in service(app).data().evidencias if e["archivo"] == "manual.log")
    assert ev["categoria"] == "Log de Qibixx"


def _upload_to(c, sid, files, tipo="Visita", registro="", categoria=""):
    return c.post(f"/visitas/{sid}/evidencias", {"tipo_registro": tipo, "registro_id": registro,
                                                "categoria": categoria, "archivo": files},
                  content_type="multipart/form-data")


def test_eliminar_evidencia_mueve_a_papelera(app):
    c = Client(app)
    svc = service(app)
    paths = app.extensions["wposs"]["paths"]
    sid = svc.create_session(SESSION_FORM)
    tid = svc.create_test(prueba_data(sid))
    _upload_to(c, sid, [(io.BytesIO(b"\xff\xd8\xffOK"), "buena.jpg"), (io.BytesIO(b"\xff\xd8\xffMAL"), "equivocada.jpg")],
               "Prueba", tid, "Foto de la máquina")
    evs = {e["archivo"]: e for e in svc.data().evidencias if e["sesion_id"] == sid}
    mala = evs[f"{tid}_maquina_equivocada.jpg"]
    # Sin confirmación explícita no se elimina.
    assert c.post(f"/carpetas/{sid}/eliminar", {"id": mala["id"]}).status_code == 400
    r = c.post(f"/carpetas/{sid}/eliminar", {"id": mala["id"], "confirmar": "si"})
    assert r.status_code == 302
    data = svc.data()
    assert [e["archivo"] for e in data.evidencias if e["sesion_id"] == sid] == [f"{tid}_maquina_buena.jpg"]
    assert not (paths.evidencias / sid / "fotos" / f"{tid}_maquina_equivocada.jpg").exists()
    papelera = list((paths.papelera / sid).iterdir())
    assert len(papelera) == 1 and papelera[0].read_bytes() == b"\xff\xd8\xffMAL"
    # La prueba asociada ya no referencia el archivo eliminado.
    test = next(t for t in data.pruebas if t["id"] == tid)
    assert test["evidencia"] == f"evidencias/{sid}/fotos/{tid}_maquina_buena.jpg"
    hist = [h for h in data.historial if h["registro_id"] == mala["id"] and h["accion"] == "Archivo eliminado"]
    assert any(h["campo"] == "Papelera" and h["nuevo"].startswith(f"papelera/{sid}/") for h in hist)
    assert c.get(f"/archivo/{mala['ruta']}").status_code == 404
    # El ID eliminado no se reutiliza.
    _upload_to(c, sid, (io.BytesIO(b"\xff\xd8\xffN"), "nueva.jpg"))
    nueva = next(e for e in svc.data().evidencias if e["archivo"] == "maquina_nueva.jpg")
    assert int(nueva["id"].split("-")[1]) > int(mala["id"].split("-")[1])


def test_eliminar_evidencia_de_otra_visita_rechazado(app):
    c = Client(app)
    svc = service(app)
    sid1 = svc.create_session(SESSION_FORM)
    sid2 = svc.create_session(SESSION_FORM)
    _upload_to(c, sid1, (io.BytesIO(b"LOG"), "a.log"))
    ev = next(e for e in svc.data().evidencias if e["sesion_id"] == sid1)
    c.post(f"/carpetas/{sid2}/eliminar", {"id": ev["id"], "confirmar": "si"})
    assert any(e["id"] == ev["id"] for e in svc.data().evidencias)
    assert (app.extensions["wposs"]["paths"].evidencias / sid1 / "logs" / "a.log").exists()


def test_eliminar_con_libro_bloqueado_queda_pendiente(workdir):
    app = make_app(workdir)
    c = Client(app)
    svc = service(app)
    paths = app.extensions["wposs"]["paths"]
    sid = svc.create_session(SESSION_FORM)
    _upload_to(c, sid, (io.BytesIO(b"LOG"), "b.log"))
    ev = next(e for e in svc.data().evidencias if e["sesion_id"] == sid)
    lock = paths.trabajo.parent / ("~$" + paths.trabajo.name[2:])
    lock.write_text("lock")
    r = c.post(f"/carpetas/{sid}/eliminar", {"id": ev["id"], "confirmar": "si"})
    assert "NO se guardó" in c.get(r.headers["Location"]).get_data(as_text=True)
    assert any(e["id"] == ev["id"] for e in svc.data().evidencias)  # no se afirma la eliminación
    lock.unlink()
    app.extensions["wposs"]["store"].retry_pending()
    assert not any(e["id"] == ev["id"] for e in svc.data().evidencias)
