"""Inicio, sesiones, IDs, edición, persistencia y datos incompletos/duplicados."""
import pytest
from openpyxl import load_workbook

from conftest import SESSION_FORM, Client, make_app, service
from wposs_vim import schema as S
from wposs_vim.excel_store import ConflictError
from wposs_vim.validators import ValidationError


def test_01_inicio_aplicacion(client):
    r = client.get("/")
    assert r.status_code == 200
    assert "Dashboard" in r.get_data(as_text=True)
    for url in ("/visitas", "/visitas/nueva", "/maquinas", "/identificacion", "/revision",
                "/pruebas", "/incidencias", "/cierre", "/reportes", "/configuracion"):
        assert client.get(url).status_code == 200, url


def test_01b_copia_de_trabajo_sin_tocar_original(app, workdir):
    paths = app.extensions["wposs"]["paths"]
    assert paths.trabajo.exists()
    assert paths.original_copia.exists()
    assert paths.trabajo != paths.original
    wb = load_workbook(paths.trabajo)
    for name in S.ORIGINAL_SHEETS + S.APP_SHEETS:
        assert name in wb.sheetnames


def test_02_crear_sesion(client, app):
    r = client.post("/visitas/nueva", {**SESSION_FORM, "accion": "guardar"})
    assert r.status_code == 302
    assert r.headers["Location"].endswith("/visitas/SES-004")
    s = service(app).get_session("SES-004")
    assert s["cliente"] == "DATAFAST" and s["comercio"] == "Comercio Piloto QA"
    assert s["pais"] == "Ecuador" and s["ciudad"] == "Quito"
    assert s["estado_registro"] == S.ESTADO_REGISTRO_EN_CURSO
    assert client.get("/visitas/SES-004").status_code == 200


def test_03_id_unico_y_no_reutilizado(client, app):
    svc = service(app)
    existing = {s["sesion_id"] for s in svc.data().sesiones}
    ids = [svc.create_session(SESSION_FORM) for _ in range(3)]
    assert ids == ["SES-004", "SES-005", "SES-006"]
    assert not existing & set(ids)
    all_ids = [s["sesion_id"] for s in svc.data().sesiones]
    assert len(all_ids) == len(set(all_ids))
    # El contador impide reutilizar números.
    assert svc.data().contadores["SES"] == "6" or int(svc.data().contadores["SES"]) == 6


def test_03b_crear_con_id_existente_no_sobrescribe(app):
    store = app.extensions["wposs"]["store"]
    before = service(app).get_session("SES-002")
    with pytest.raises(ConflictError):
        store.write([("create_session", {"sesion": {**SESSION_FORM, "sesion_id": "SES-002"},
                                         "estado_registro": "En curso"})], "duplicado")
    assert service(app).get_session("SES-002")["cliente"] == before["cliente"]


def test_04_editar_sesion_con_historial(client, app):
    sid = service(app).create_session(SESSION_FORM)
    r = client.post(f"/visitas/{sid}/editar", {**SESSION_FORM, "firmware": "13.0.17", "accion": "guardar"})
    assert r.status_code == 302
    s = service(app).get_session(sid)
    assert s["firmware"] == "13.0.17"
    hist = [h for h in service(app).data().historial if h["registro_id"] == sid and h["campo"] == "Firmware NEW5320"]
    assert hist and hist[-1]["anterior"] == "13.0.22" and hist[-1]["nuevo"] == "13.0.17"


def test_04b_borrador_y_continuar(client, app):
    r = client.post("/visitas/nueva", {"comercio": "Solo comercio", "accion": "borrador"})
    assert r.status_code == 302
    sid = r.headers["Location"].rsplit("/", 1)[-1]
    assert service(app).get_session(sid)["estado_registro"] == S.ESTADO_REGISTRO_BORRADOR
    # Completar el borrador lo pasa a "En curso".
    client.post(f"/visitas/{sid}/editar", {**SESSION_FORM, "accion": "guardar"})
    assert service(app).get_session(sid)["estado_registro"] == S.ESTADO_REGISTRO_EN_CURSO


def test_10_persistencia_tras_reinicio(workdir):
    app1 = make_app(workdir)
    sid = service(app1).create_session(SESSION_FORM)
    tid = service(app1).create_test({"sesion_id": sid, "componente": "POS NEW5320",
                                     "tipo_verificacion": "Prueba ejecutada", "objetivo": "Encendido",
                                     "estado": "Pendiente"})
    del app1
    app2 = make_app(workdir)
    s = service(app2).get_session(sid)
    assert s is not None and s["serial_pos"] == "POS-SERIAL-1"
    assert any(t["id"] == tid for t in service(app2).data().pruebas)
    # Después de reiniciar el siguiente ID sigue siendo único.
    assert service(app2).next_id("SES") == "SES-005"


def test_15_datos_incompletos_conservan_formulario(client, app):
    before = len(service(app).data().sesiones)
    form = {**SESSION_FORM, "cliente": "", "fabricante": "", "fecha": "27/10/2026", "accion": "guardar"}
    r = client.post("/visitas/nueva", form)
    assert r.status_code == 422
    html = r.get_data(as_text=True)
    assert "Campo obligatorio." in html
    assert "Use el formato AAAA-MM-DD." in html
    assert 'value="Validar MDB"' not in html  # objetivo es textarea
    assert "Validar MDB" in html and '<option value="13.0.22" selected>' in html  # datos conservados
    assert len(service(app).data().sesiones) == before


def test_15b_catalogo_invalido_rechazado(app):
    with pytest.raises(ValidationError) as exc:
        service(app).create_session({**SESSION_FORM, "tipo_prueba": "Inventado"})
    assert "tipo_prueba" in exc.value.errors


def test_15c_texto_con_igual_no_se_guarda_como_formula(app):
    sid = service(app).create_session({**SESSION_FORM, "observaciones": "=HYPERLINK(\"http://x\")"})
    paths = app.extensions["wposs"]["paths"]
    wb = load_workbook(paths.trabajo)
    ws = wb[S.SHEET_ID]
    row = next(r for r in range(2, ws.max_row + 1) if ws.cell(r, 1).value == sid)
    cell = ws.cell(row, 26)
    assert cell.data_type == "s"


def test_15d_fabricante_nuevo_amplia_catalogo_sin_duplicar(app):
    svc = service(app)
    svc.create_session({**SESSION_FORM, "fabricante": "Nuevo Fab"})
    svc.create_session({**SESSION_FORM, "fabricante": "Nuevo Fab"})
    assert svc.data().catalogos["fabricante"].count("Nuevo Fab") == 1


def test_seguridad_csrf_y_acceso_local(app):
    c = app.test_client()
    assert c.post("/visitas/nueva", data={"cliente": "x"}).status_code == 400
    r = c.get("/", environ_base={"REMOTE_ADDR": "192.168.1.20"})
    assert r.status_code == 403


def test_ruta_manipulada_no_sale_de_evidencias(client):
    assert client.get("/archivo/../trabajo/WPOSS_Vending_Datos.xlsx").status_code == 404
    assert client.get("/archivo/evidencias/..%2F..%2Fconfig.json").status_code == 404
    assert client.get("/archivo/trabajo/WPOSS_Vending_Datos.xlsx").status_code == 404
    assert client.get("/visitas/..%2Fx").status_code == 404
