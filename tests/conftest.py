import shutil
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from wposs_vim import create_app  # noqa: E402

# Copia anonimizada del libro WPOSS (misma estructura y fórmulas, sin datos de clientes).
FIXTURE = ROOT / "tests" / "fixtures" / "checklist_anonimizado.xlsx"
ORIGINAL_NAME = "WPOSS_Checklist_Pruebas_Vending_5320.xlsx"


@pytest.fixture
def workdir(tmp_path):
    """Directorio aislado con una copia del libro original."""
    shutil.copy2(FIXTURE, tmp_path / ORIGINAL_NAME)
    return tmp_path


@pytest.fixture
def empty_workdir(tmp_path):
    """Libro original sin ningún registro de datos (solo estructura)."""
    target = tmp_path / ORIGINAL_NAME
    wb = load_workbook(FIXTURE)
    for name in ("Identificación y Versiones", "Revisión Física y Eléctrica", "Casos de Prueba", "Incidencias"):
        ws = wb[name]
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.value = None
    for r in range(2, wb["Cierre de Sesión"].max_row + 1):
        wb["Cierre de Sesión"].cell(r, 1).value = None
    wb.save(target)
    return tmp_path


def make_app(base, **overrides):
    return create_app(base, config_overrides=overrides or None, testing=True)


@pytest.fixture
def app(workdir):
    return make_app(workdir)


class Client:
    """Cliente de pruebas que agrega el token CSRF a los POST."""

    def __init__(self, app):
        self.app = app
        self.c = app.test_client()
        self.c.get("/")
        with self.c.session_transaction() as s:
            self.token = s["_csrf"]

    def get(self, url, **kw):
        return self.c.get(url, **kw)

    def post(self, url, data=None, **kw):
        data = dict(data or {})
        data["_csrf"] = self.token
        return self.c.post(url, data=data, **kw)


@pytest.fixture
def client(app):
    return Client(app)


def service(app):
    return app.extensions["wposs"]["service"]


SESSION_FORM = {
    "fecha": "2026-10-27", "hora": "09:15", "responsable": "Técnico Prueba",
    "participantes": "Equipo WPOSS", "cliente": "Cliente Piloto QA", "pais": "Colombia",
    "ciudad": "Bogotá", "tipo_prueba": "Compatibilidad", "objetivo": "Validar MDB",
    "fabricante": "Fabricante QA", "modelo_maquina": "Modelo QA-1", "serial_maquina": "SM-001",
    "modelo_pos": "NEW5320", "serial_pos": "POS-SERIAL-1", "firmware": "13.0.30",
    "mdb_estandar": "MDB 4.2", "mdb_version": "1.0", "mdb_modulo": "mdb_mod",
    "version_financiera": "FIN 2.0", "version_polaris": "AGP 1.1",
}


def prueba_data(sesion_id, **kw):
    data = {"sesion_id": sesion_id, "componente": "MDB", "tipo_verificacion": "Prueba ejecutada",
            "objetivo": "Verificar dirección cashless", "resultado_esperado": "Dirección correcta",
            "estado": "Pendiente", "responsable": "QA"}
    data.update(kw)
    return data


