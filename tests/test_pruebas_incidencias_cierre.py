"""Pruebas, incidencias, asociación, revisión física, cierre y resumen."""
import pytest

from conftest import SESSION_FORM, prueba_data, service
from wposs_vim import schema as S
from wposs_vim.validators import ValidationError


@pytest.fixture
def sid(app):
    return service(app).create_session(SESSION_FORM)


def test_05_multiples_pruebas_por_sesion(client, app, sid):
    ids = []
    for comp in ("MDB", "MDB", "Financiera", "POS NEW5320"):
        r = client.post("/pruebas/nueva", prueba_data(sid, componente=comp))
        assert r.status_code == 302
    tests = [t for t in service(app).data().pruebas if t["sesion_id"] == sid]
    ids = [t["id"] for t in tests]
    assert len(tests) == 4
    assert len(set(ids)) == 4
    # Prefijos según el componente, numeración sin colisión con el libro original.
    assert ids[0] == "MDB-004" and ids[1] == "MDB-005"
    assert ids[2].startswith("FIN-") and ids[3].startswith("POS-")
    # Las pruebas existentes del libro se conservan.
    assert any(t["id"] == "MDB-002" for t in service(app).data().pruebas)


def test_05b_no_aprobar_sin_resultado(app, sid):
    with pytest.raises(ValidationError) as exc:
        service(app).create_test(prueba_data(sid, estado="Aprobado"))
    assert "resultado_obtenido" in exc.value.errors
    with pytest.raises(ValidationError) as exc:
        service(app).create_test(prueba_data(sid, estado="Fallido", resultado_esperado="",
                                             resultado_obtenido="x"))
    assert "resultado_esperado" in exc.value.errors
    # Guardar una prueba nunca la aprueba automáticamente.
    tid = service(app).create_test(prueba_data(sid))
    t = next(t for t in service(app).data().pruebas if t["id"] == tid)
    assert t["estado"] == "Pendiente"


def test_06_registrar_incidencia(client, app, sid):
    r = client.post("/incidencias/nueva", {"sesion_id": sid, "fecha": "2026-10-27",
                                           "descripcion": "Sin respuesta al POLL", "severidad": "Alta",
                                           "estado": "Abierto", "responsable": "Desarrollo"})
    assert r.status_code == 302
    incs = [i for i in service(app).data().incidencias if i["sesion_id"] == sid]
    assert len(incs) == 1 and incs[0]["id"] == "INC-004"


def test_07_asociacion_incidencia_prueba(client, app, sid):
    svc = service(app)
    tid = svc.create_test(prueba_data(sid, estado="Fallido", resultado_obtenido="Sin respuesta"))
    defaults = svc.incident_defaults_from_test(tid)
    assert defaults["sesion_id"] == sid and defaults["prueba_id"] == tid
    assert defaults["componente_version"] == "MDB 1.0"
    r = client.get(f"/incidencias/nueva?prueba={tid}")
    assert sid in r.get_data(as_text=True)
    form = {**defaults, "severidad": "Crítica"}
    r = client.post("/incidencias/nueva", form)
    assert r.status_code == 302
    data = svc.data()
    inc = next(i for i in data.incidencias if i["sesion_id"] == sid)
    test = next(t for t in data.pruebas if t["id"] == tid)
    assert inc["prueba_id"] == tid
    assert test["incidencia"] == inc["id"]
    # Modificar la prueba (aprobarla) no cierra la incidencia.
    svc.update_test(tid, {**prueba_data(sid), "estado": "Aprobado", "resultado_obtenido": "OK en reprueba"})
    inc2 = next(i for i in svc.data().incidencias if i["id"] == inc["id"])
    assert inc2["estado"] == "Abierto"
    # Cambio de estado manual queda en el historial.
    svc.update_incident(inc["id"], {**inc2, "estado": "Cerrado", "resultado_reprueba": "OK"})
    hist = svc.history("Incidencia", inc["id"])
    assert any("Cambio de estado" in h["accion"] and h["nuevo"] == "Cerrado" for h in hist)


def test_07b_prueba_de_otra_sesion_rechazada(app, sid):
    with pytest.raises(ValidationError):
        service(app).create_incident({"sesion_id": sid, "fecha": "2026-10-27", "descripcion": "x",
                                      "severidad": "Baja", "estado": "Abierto", "prueba_id": "MDB-002"})


def test_revision_fisica_sin_duplicar_ni_alterar_historico(client, app, sid):
    svc = service(app)
    before_other = [r for r in svc.data().revisiones if r["sesion_id"] == "SES-003"]
    client.post(f"/visitas/{sid}/revision/iniciar")
    client.post(f"/visitas/{sid}/revision/iniciar")  # segunda vez no duplica
    rows = [r for r in svc.data().revisiones if r["sesion_id"] == sid]
    assert len(rows) == 8 and all(r["resultado"] == "Pendiente" for r in rows)
    form = {}
    for i, r in enumerate(rows):
        form[f"item_{i}"] = r["item"]
        form[f"resultado_{i}"] = "Aprobado"
        form[f"observaciones_{i}"] = ""
    form["resultado_1"] = "Fallido"
    assert client.post(f"/visitas/{sid}/revision", form).status_code == 302
    rows = [r for r in svc.data().revisiones if r["sesion_id"] == sid]
    assert rows[1]["resultado"] == "Fallido"
    assert [r for r in svc.data().revisiones if r["sesion_id"] == "SES-003"] == before_other


def test_08_totales_del_cierre(app, sid):
    svc = service(app)
    for estado, obtenido in (("Aprobado", "ok"), ("Aprobado", "ok"), ("Fallido", "error"),
                             ("Bloqueado", ""), ("Pendiente", ""), ("No aplica", "")):
        svc.create_test(prueba_data(sid, estado=estado, resultado_obtenido=obtenido))
    svc.create_incident({"sesion_id": sid, "fecha": "2026-10-27", "descripcion": "a",
                         "severidad": "Media", "estado": "Abierto"})
    svc.create_incident({"sesion_id": sid, "fecha": "2026-10-27", "descripcion": "b",
                         "severidad": "Media", "estado": "Cerrado"})
    t = svc.totals(sid)
    assert t["planificadas"] == 6
    assert t["ejecutadas"] == 5  # fórmula: distinto de "Pendiente" y de vacío
    assert (t["aprobadas"], t["fallidas"], t["bloqueadas"], t["pendientes"]) == (2, 1, 1, 1)
    assert t["incidencias_abiertas"] == 1
    ev = svc.evaluate_closure(sid)
    assert not ev["puede_apto"] and ev["sugerido"] == S.CONCEPTO_NO_APTO
    # No se puede declarar apto aunque la mayoría pase.
    with pytest.raises(ValidationError) as exc:
        svc.save_cierre(sid, {"fecha_cierre": "2026-10-28", "concepto": S.CONCEPTO_APTO,
                              "responsable_cierre": "QA"})
    assert "concepto" in exc.value.errors
    # "Apto con condiciones" exige documentar los pendientes.
    with pytest.raises(ValidationError):
        svc.save_cierre(sid, {"fecha_cierre": "2026-10-28", "concepto": S.CONCEPTO_CONDICIONES,
                              "responsable_cierre": "QA", "pendientes_area": ""})
    svc.save_cierre(sid, {"fecha_cierre": "2026-10-28", "concepto": S.CONCEPTO_NO_APTO,
                          "responsable_cierre": "QA", "pendientes_area": "Desarrollo: MDB"})
    s = svc.get_session(sid)
    assert s["concepto"] == S.CONCEPTO_NO_APTO and s["estado_registro"] == S.ESTADO_REGISTRO_CERRADA


def test_08b_cierre_escribe_formulas_originales(app, sid):
    from openpyxl import load_workbook
    svc = service(app)
    svc.create_test(prueba_data(sid, estado="Aprobado", resultado_obtenido="ok"))
    svc.save_cierre(sid, {"fecha_cierre": "2026-10-28", "concepto": S.CONCEPTO_PENDIENTE,
                          "responsable_cierre": "QA"})
    wb = load_workbook(app.extensions["wposs"]["paths"].trabajo)
    ws = wb[S.SHEET_CIERRE]
    row = next(r for r in range(2, ws.max_row + 1) if ws.cell(r, 1).value == sid)
    assert row == 5  # primera fila preparada libre
    assert ws.cell(row, 19).value == f"=IF(A{row}=\"\",\"\",COUNTIF('Casos de Prueba'!$B$2:$B$1000,A{row}))"
    assert ws.cell(row, 30).value.startswith(f"=IF(A{row}=\"\",\"\",\"Resumen de visita")
    assert ws.cell(row, 28).value == S.CONCEPTO_PENDIENTE


def test_08c_apto_permitido_cuando_cumple_criterios(app, sid):
    svc = service(app)
    svc.init_revision(sid)
    rows = [r for r in svc.data().revisiones if r["sesion_id"] == sid]
    form = {}
    for i, r in enumerate(rows):
        form.update({f"item_{i}": r["item"], f"resultado_{i}": "Aprobado", f"observaciones_{i}": ""})
    svc.update_revision(sid, form)
    svc.create_test(prueba_data(sid, estado="Aprobado", resultado_obtenido="ok"))
    assert svc.evaluate_closure(sid)["puede_apto"]
    svc.save_cierre(sid, {"fecha_cierre": "2026-10-28", "concepto": S.CONCEPTO_APTO, "responsable_cierre": "QA"})
    assert svc.get_session(sid)["categoria"] == "aprobada"


def test_08d_componente_critico_configurable(workdir):
    from conftest import make_app
    app = make_app(workdir, cierre={"componentes_criticos": ["Financiera"]})
    svc = service(app)
    sid = svc.create_session(SESSION_FORM)
    svc.create_test(prueba_data(sid, estado="Aprobado", resultado_obtenido="ok"))
    assert any("Financiera" in b for b in svc.evaluate_closure(sid)["bloqueos"])


def test_09_resumen_para_correo(client, app, sid):
    svc = service(app)
    svc.create_test(prueba_data(sid, estado="Aprobado", resultado_obtenido="ok"))
    svc.create_test(prueba_data(sid))
    svc.save_cierre(sid, {"fecha_cierre": "2026-10-28", "concepto": S.CONCEPTO_PENDIENTE,
                          "responsable_cierre": "QA", "pendientes_area": "QA: MDB-005"})
    text = client.get(f"/visitas/{sid}/resumen.txt").get_data(as_text=True)
    assert text.startswith(f"Resumen de visita — Sesión {sid}\n")
    assert "Máquina: Fabricante QA Modelo QA-1 (Serial: SM-001)" in text
    assert "POS: NEW5320 (Serial: POS-SERIAL-1) — Firmware: 13.0.22" in text
    assert "MDB: MDB 4.2 / 1.0 / mdb_mod" in text
    assert "Resultado de pruebas: 1/2 ejecutadas — Aprobadas: 1, Fallidas: 0, Bloqueadas: 0, Pendientes: 1" in text
    assert "Concepto final: Pendiente" in text
    assert text.rstrip().endswith("Pendientes: QA: MDB-005")
    page = client.get(f"/visitas/{sid}/cierre").get_data(as_text=True)
    assert "Resumen para correo" in page


def test_ejemplos_excluidos_de_totales(app):
    svc = service(app)
    # INC-002 es una fila de ejemplo ("Fila de ejemplo — reemplazar...") en estado Abierto.
    assert svc.totals("SES-002")["incidencias_abiertas"] == 0
    assert svc.totals("SES-002", include_examples=True)["incidencias_abiertas"] == 1
    assert all(s["sesion_id"] != "SES-001" for s in svc.sessions())


def test_eliminar_fila_de_ejemplo_requiere_marca(client, app):
    r = client.post("/registros/eliminar-ejemplo", {"tipo": "Incidencia", "id": "INC-002",
                                                     "sesion_id": "SES-002", "confirmar": "si"})
    assert r.status_code == 302
    assert all(i["id"] != "INC-002" for i in service(app).data().incidencias)
    assert any(h["registro_id"] == "INC-002" for h in service(app).data().historial)
    # Una incidencia real no puede eliminarse por esta vía.
    sid = service(app).create_session(SESSION_FORM)
    iid = service(app).create_incident({"sesion_id": sid, "fecha": "2026-10-27", "descripcion": "real",
                                        "severidad": "Baja", "estado": "Abierto"})
    client.post("/registros/eliminar-ejemplo", {"tipo": "Incidencia", "id": iid, "sesion_id": sid,
                                                "confirmar": "si"})
    assert any(i["id"] == iid for i in service(app).data().incidencias)
